import csv
import json
from pathlib import Path

import grpc
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import Float, Integer, String, UniqueConstraint, select
from sqlalchemy.orm import Mapped, Session, mapped_column, sessionmaker

from proto_generated import academics_pb2, academics_pb2_grpc
from proto_generated import grades_pb2, grades_pb2_grpc
from proto_generated import periods_pb2, periods_pb2_grpc
from shared.app_common.auth import require_roles
from shared.app_common.config import BaseServiceSettings
from shared.app_common.database import Base, create_session_factory, session_scope
from shared.app_common.files import resolve_input_file
from shared.app_common.grpc_runtime import start_grpc_server
from shared.app_common.responses import ok


class Settings(BaseServiceSettings):
    app_name: str = "AGM Calificaciones & Ponderaciones"
    service_slug: str = "ms-grades"
    rest_port: int = 8014
    grpc_port: int = 50054
    # ¡ADIÓS SQLITE! Apuntamos a la base de datos exclusiva de Calificaciones
    database_url: str = "postgresql+psycopg://agm:agm_dev_password@postgres:5432/agm_grades_db"
    
    # Rutas internas para comunicarse con otros microservicios
    periods_grpc_target: str = "ms-periods:50052"
    academics_grpc_target: str = "ms-academics:50053"


class WeightCategory(Base):
    __tablename__ = "weight_categories"

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    materia_id: Mapped[int] = mapped_column(Integer, index=True)
    nombre: Mapped[str] = mapped_column(String(120))
    porcentaje: Mapped[float] = mapped_column(Float)


class Activity(Base):
    __tablename__ = "activities"

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    materia_id: Mapped[int] = mapped_column(Integer, index=True)
    categoria_id: Mapped[int] = mapped_column(Integer, index=True)
    nombre: Mapped[str] = mapped_column(String(150))
    max_puntos: Mapped[float] = mapped_column(Float, default=100.0)


class Grade(Base):
    __tablename__ = "grades"
    __table_args__ = (UniqueConstraint("activity_id", "student_id", name="uq_activity_student"),)

    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    activity_id: Mapped[int] = mapped_column(Integer, index=True)
    student_id: Mapped[int] = mapped_column(Integer, index=True)
    score: Mapped[float] = mapped_column(Float)


class WeightItem(BaseModel):
    nombre: str
    porcentaje: float = Field(gt=0, le=100)


class WeightPayload(BaseModel):
    items: list[WeightItem]

class ActivityUpdatePayload(BaseModel):
    categoria_id: int
    nombre: str
    max_puntos: float = Field(default=100, gt=0)

class ActivityPayload(BaseModel):
    materia_id: int
    categoria_id: int
    nombre: str
    max_puntos: float = Field(default=100, gt=0)


class GradePayload(BaseModel):
    activity_id: int
    student_id: int
    score: float = Field(ge=0)


def get_students_for_subject(target: str, materia_id: int) -> list[academics_pb2.AlumnoInfo]:
    with grpc.insecure_channel(target) as channel:
        stub = academics_pb2_grpc.AcademicsServiceStub(channel)
        reply = stub.GetAlumnosByMateria(academics_pb2.MateriaIdRequest(materia_id=materia_id), timeout=3)
        return list(reply.items)


def ensure_subject(target: str, materia_id: int) -> None:
    try:
        with grpc.insecure_channel(target) as channel:
            stub = periods_pb2_grpc.PeriodsServiceStub(channel)
            materia = stub.GetMateriaById(periods_pb2.MateriaIdRequest(materia_id=materia_id))
            if not materia.materia_id:
                raise HTTPException(status_code=404, detail="Materia no encontrada")
    except grpc.RpcError as exc:
        raise HTTPException(status_code=503, detail="Periods service unavailable") from exc


def category_dict(item: WeightCategory) -> dict:
    return {"id": item.id, "materia_id": item.materia_id, "nombre": item.nombre, "porcentaje": item.porcentaje}


def compute_student_average(session: Session, materia_id: int, student_id: int) -> tuple[float, int]:
    categories = session.scalars(select(WeightCategory).where(WeightCategory.materia_id == materia_id)).all()
    total = 0.0
    for category in categories:
        activities = session.scalars(select(Activity).where(Activity.categoria_id == category.id)).all()
        if not activities:
            continue
        category_score = 0.0
        for activity in activities:
            grade = session.scalar(
                select(Grade).where(Grade.activity_id == activity.id, Grade.student_id == student_id)
            )
            if grade is None:
                continue
            category_score += (grade.score / activity.max_puntos) * 100
        category_score = category_score / len(activities)
        total += category_score * (category.porcentaje / 100)
    rounded = int(total + 0.5)
    return round(total, 2), rounded


def parse_grade_import(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            return [
                {"matricula": str(row.get("matricula") or "").strip(), "score": float(row.get("calificacion") or 0)}
                for row in reader
                if str(row.get("matricula") or "").strip()
            ]
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.active
    header = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        matricula = str(row.get("matricula") or "").strip()
        if not matricula:
            continue
        rows.append({"matricula": matricula, "score": float(row.get("calificacion") or 0)})
    return rows


class GradesGrpcService(grades_pb2_grpc.GradesServiceServicer):
    def __init__(self, session_factory: sessionmaker[Session], settings: Settings):
        self.session_factory = session_factory
        self.settings = settings

    def GetConcentrado(self, request, context):
        with session_scope(self.session_factory) as session:
            try:
                students = get_students_for_subject(self.settings.academics_grpc_target, request.materia_id)
            except grpc.RpcError:
                students = []
            items = []
            for student in students:
                promedio_real, promedio_redondeado = compute_student_average(session, request.materia_id, student.alumno_id)
                items.append(
                    grades_pb2.AlumnoCalif(
                        alumno_id=student.alumno_id,
                        nombre=student.nombre,
                        promedio_real=promedio_real,
                        promedio_redondeado=promedio_redondeado,
                    )
                )
            return grades_pb2.ConcentradoReply(items=items)

    def GetPromedioAlumno(self, request, context):
        with session_scope(self.session_factory) as session:
            promedio_real, _ = compute_student_average(session, request.materia_id, request.alumno_id)
            return grades_pb2.FloatReply(value=promedio_real)

    def GetEstadisticasMateria(self, request, context):
        with session_scope(self.session_factory) as session:
            students = get_students_for_subject(self.settings.academics_grpc_target, request.materia_id)
            if not students:
                return grades_pb2.StatsReply(total_alumnos=0, promedio_grupal=0, aprobacion=0)
            averages = [compute_student_average(session, request.materia_id, item.alumno_id)[0] for item in students]
            aprobados = len([value for value in averages if value >= 70])
            return grades_pb2.StatsReply(
                total_alumnos=len(averages),
                promedio_grupal=round(sum(averages) / len(averages), 2),
                aprobacion=round((aprobados / len(averages)) * 100, 2),
            )


app = FastAPI(title="AGM Calificaciones & Ponderaciones", version="0.1.0", root_path="/api/grades")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup_event() -> None:
    settings = Settings()
    session_factory = create_session_factory(settings.database_url)
    Base.metadata.create_all(session_factory.kw["bind"])
    app.state.settings = settings
    app.state.session_factory = session_factory
    app.state.grpc_server, app.state.grpc_thread = start_grpc_server(
        settings.grpc_port,
        lambda server: grades_pb2_grpc.add_GradesServiceServicer_to_server(GradesGrpcService(session_factory, settings), server),
    )


@app.on_event("shutdown")
def shutdown_event() -> None:
    grpc_server = getattr(app.state, "grpc_server", None)
    if grpc_server is not None:
        grpc_server.stop(grace=1)


@app.get("/health")
def health() -> dict:
    return ok({"service": "grades", "status": "ok"})


@app.get("/ponderaciones/{materia_id}")
def list_weights(materia_id: int, request: Request, user=Depends(require_roles("admin", "docente", "alumno"))) -> dict:
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        items = session.scalars(select(WeightCategory).where(WeightCategory.materia_id == materia_id)).all()
        return ok([category_dict(item) for item in items])


@app.post("/ponderaciones/{materia_id}")
@app.put("/ponderaciones/{materia_id}")
def save_weights(materia_id: int, payload: WeightPayload, request: Request, user=Depends(require_roles("admin", "docente"))) -> dict:
    total = round(sum(item.porcentaje for item in payload.items), 2)
    if total != 100:
        raise HTTPException(status_code=400, detail="La suma de ponderaciones debe ser exactamente 100")
    ensure_subject(request.app.state.settings.periods_grpc_target, materia_id)
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        for existing in session.scalars(select(WeightCategory).where(WeightCategory.materia_id == materia_id)).all():
            session.delete(existing)
        for item in payload.items:
            session.add(WeightCategory(materia_id=materia_id, nombre=item.nombre, porcentaje=item.porcentaje))
        return ok({"materia_id": materia_id, "total": total}, "Ponderaciones guardadas")


@app.post("/actividades")
def create_activity(payload: ActivityPayload, request: Request, user=Depends(require_roles("admin", "docente"))) -> dict:
    ensure_subject(request.app.state.settings.periods_grpc_target, payload.materia_id)
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        category = session.get(WeightCategory, payload.categoria_id)
        if category is None:
            raise HTTPException(status_code=404, detail="Categoría no encontrada")
        activity = Activity(**payload.model_dump())
        session.add(activity)
        session.flush()
        return ok({"id": activity.id}, "Actividad creada")


@app.get("/actividades/materia/{materia_id}")
def list_activities(materia_id: int, request: Request, user=Depends(require_roles("admin", "docente", "alumno"))) -> dict:
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        # Recuperamos todas las actividades de la materia solicitada
        activities = session.scalars(select(Activity).where(Activity.materia_id == materia_id)).all()
        
        # Devolvemos una lista de diccionarios limpios al frontend
        return ok([
            {
                "id": act.id,
                "materia_id": act.materia_id,
                "categoria_id": act.categoria_id,
                "nombre": act.nombre,
                "max_puntos": act.max_puntos
            }
            for act in activities
        ])




@app.post("/calificaciones")
def upsert_grade(payload: GradePayload, request: Request, user=Depends(require_roles("admin", "docente"))) -> dict:
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        activity = session.get(Activity, payload.activity_id)
        if activity is None:
            raise HTTPException(status_code=404, detail="Actividad no encontrada")
        grade = session.scalar(select(Grade).where(Grade.activity_id == payload.activity_id, Grade.student_id == payload.student_id))
        if grade is None:
            grade = Grade(**payload.model_dump())
            session.add(grade)
        else:
            grade.score = payload.score
        return ok(None, "Calificación guardada")


@app.post("/calificaciones/importar")
async def import_grades(
    request: Request,
    activity_id: int = Form(...),
    source_path: str | None = Form(None),
    file: UploadFile | None = File(None),
    user=Depends(require_roles("admin", "docente")),
) -> dict:
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        activity = session.get(Activity, activity_id)
        if activity is None:
            raise HTTPException(status_code=404, detail="Actividad no encontrada")
        students = {item.matricula: item.alumno_id for item in get_students_for_subject(request.app.state.settings.academics_grpc_target, activity.materia_id)}
        path = await resolve_input_file(upload=file, source_path=source_path)
        rows = parse_grade_import(path)
        updated = 0
        for row in rows:
            student_id = students.get(row["matricula"])
            if student_id is None:
                continue
            grade = session.scalar(select(Grade).where(Grade.activity_id == activity_id, Grade.student_id == student_id))
            if grade is None:
                grade = Grade(activity_id=activity_id, student_id=student_id, score=row["score"])
                session.add(grade)
            else:
                grade.score = row["score"]
            updated += 1
        return ok({"filas": len(rows), "actualizadas": updated}, "Calificaciones importadas")


@app.get("/concentrado/{materia_id}")
def get_concentrado(materia_id: int, request: Request, user=Depends(require_roles("admin", "docente", "alumno"))) -> dict:
    session_factory = request.app.state.session_factory
    with session_scope(session_factory) as session:
        students = get_students_for_subject(request.app.state.settings.academics_grpc_target, materia_id)
        rows = []
        for student in students:
            real, rounded = compute_student_average(session, materia_id, student.alumno_id)
            rows.append(
                {
                    "alumno_id": student.alumno_id,
                    "matricula": student.matricula,
                    "nombre": student.nombre,
                    "promedio_real": real,
                    "promedio_redondeado": rounded,
                }
            )
        return ok(rows)
